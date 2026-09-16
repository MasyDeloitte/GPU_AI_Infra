from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "AI_Sizing_infraSizing_industry_Baseline_v6.2.xlsx"
TARGET = ROOT / "Forge_Catalog_Reference_v1.xlsx"
CATALOG_SHEETS = [
    "Model_Mapping_Catalog",
    "Model_Selection",
    "Hardware_Mapping_Catalog",
    "GPU_Server_Selection",
]
MIN_DATA_ROWS = 30


def copy_sheet_data(source_sheet, target_sheet):
    for row in source_sheet.iter_rows():
        for source_cell in row:
            target_cell = target_sheet[source_cell.coordinate]
            target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.number_format = source_cell.number_format
                target_cell.protection = copy(source_cell.protection)
    for key, dimension in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[key].width = dimension.width
    for key, dimension in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[key].height = dimension.height
    for merged in source_sheet.merged_cells.ranges:
        target_sheet.merge_cells(str(merged))
    target_sheet.freeze_panes = source_sheet.freeze_panes
    target_sheet.auto_filter.ref = source_sheet.auto_filter.ref


def data_start(sheet):
    return 4 if sheet.title != "Model_Selection" else 4


def last_data_row(sheet):
    last = 3
    for row in range(4, sheet.max_row + 1):
        if any(sheet.cell(row, col).value is not None for col in range(1, sheet.max_column + 1)):
            last = row
    return last


def add_template_rows(sheet):
    last = last_data_row(sheet)
    existing = max(0, last - 3)
    needed = max(0, MIN_DATA_ROWS - existing)
    if needed == 0:
        return 0
    header_values = [sheet.cell(3, col).value for col in range(1, sheet.max_column + 1)]
    for offset in range(1, needed + 1):
        row = last + offset
        for col, header in enumerate(header_values, start=1):
            if col == 1:
                value = f"USER-DEFINED-{sheet.title[:3].upper()}-{existing + offset:03d}"
            elif col == 2:
                value = "User-defined catalog entry"
            elif col == 3:
                value = "Data Required"
            elif header in {"Status", "Value Status", "Readiness", "Benchmark Type", "Selection Status"}:
                value = "Review Required"
            elif header in {"Source", "Source Key", "Source / Formula", "Production Replacement", "Notes"}:
                value = "Add source and validation evidence"
            else:
                value = "Data Required"
            cell = sheet.cell(row, col, value)
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.font = Font(color="7F6000")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row].height = 36
    return needed


def main():
    source = openpyxl.load_workbook(SOURCE, data_only=False)
    workbook = Workbook()
    workbook.remove(workbook.active)

    index = workbook.create_sheet("Catalog_Index")
    index.append(["Forge Catalog Reference Workbook", None, None, None])
    index.append(["Catalog Version", date.today().isoformat(), None, None])
    index.append(["Source Workbook", SOURCE.name, None, None])
    index.append(["Purpose", "Human-maintained catalog source for the web calculator", None, None])
    index.append([])
    index.append(["Catalog Sheet", "Rows Included", "Minimum Rows", "Maintenance Rule"])

    for sheet_name in CATALOG_SHEETS:
        target = workbook.create_sheet(sheet_name)
        copy_sheet_data(source[sheet_name], target)
        added = add_template_rows(target)
        target.sheet_view.showGridLines = False
        target.freeze_panes = "C4"
        index.append([sheet_name, max(30, last_data_row(target) - 3), MIN_DATA_ROWS, "Replace yellow USER-DEFINED rows only with reviewed source-backed data"])
        target["A1"] = f"InfraPilot AI Catalog Reference | {sheet_name} | Version {date.today().isoformat()}"
        target["A1"].font = Font(bold=True, color="FFFFFF", size=13)
        target["A1"].fill = PatternFill("solid", fgColor="17365D")
        target.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(6, target.max_column))
        for cell in target[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col in range(1, target.max_column + 1):
            letter = get_column_letter(col)
            if target.column_dimensions[letter].width is None:
                target.column_dimensions[letter].width = 18

    index["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    index["A1"].fill = PatternFill("solid", fgColor="17365D")
    index.merge_cells("A1:D1")
    for cell in index[6]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    index.column_dimensions["A"].width = 32
    index.column_dimensions["B"].width = 18
    index.column_dimensions["C"].width = 16
    index.column_dimensions["D"].width = 72
    index.freeze_panes = "A7"
    index.sheet_view.showGridLines = False

    workbook.save(TARGET)
    print(f"Created {TARGET}")


if __name__ == "__main__":
    main()
